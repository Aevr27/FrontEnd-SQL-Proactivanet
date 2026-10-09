"""Historial de fechas de iniciativas desde los comentarios del Excel de Experiencia.

Lee los threaded comments de DBProblems en Fecha Analisis / Fecha Solucion /
Fecha Cierre, rearma la cadena de cada celda y deja:

  <salida>/ProblemFechaEvento_Revision.csv   una fila por fecha comentada, con
                                             su decision (CUENTA / NO_CUENTA /
                                             EXCLUIDA_CIERRE) y sus banderas
  <salida>/ProblemFechaEvento_Diferencias.csv  solo con --anterior: que cambio
                                             contra una revision anterior
  <salida>/ProblemFechaEvento_Importar.sql   staging + vista previa + INSERT en
                                             transaccion; NO cambia nada salvo
                                             que se pongan @Confirmar = 1 y
                                             @ModoFechaRegistro a mano

No se conecta a ninguna base. Las salidas llevan datos del Excel (codigos,
autores de comentarios): se quedan en el equipo, nunca en el repositorio.

Reglas (fijadas 2026-10-08):
  - texto del comentario = fecha ANTERIOR; dT del comentario = momento
    aproximado del cambio; la fecha NUEVA es la del siguiente comentario con
    fecha, o el valor actual de la celda para el ultimo;
  - Analisis y Solucion: cuenta todo cambio real (anterior != nueva),
    adelantos y reversiones incluidos;
  - Cierre: solo extensiones (nueva > anterior).

Uso:
  python -I extraer_historial_fechas.py <libro.xlsx> <carpeta_salida> [--anterior <Revision.csv>]
"""
import argparse
import collections
import csv
import datetime
import os
import re
import sys
import unicodedata
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
from openpyxl.utils import column_index_from_string, coordinate_to_tuple, get_column_letter

HOJA = 'DBProblems'
NS = {
    'm': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'tc': 'http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments',
}

# Encabezado normalizado (sin acentos, espacios ni mayusculas) -> papel.
# Las columnas se buscan por nombre; si falta alguna, se para.
ENCABEZADOS = {
    'codigo': 'codigo',
    'fechaanalisis': 'FechaAnalisis',
    'nrocambiofechaanalisis': 'cnt:FechaAnalisis',
    'foriginalsolucion': 'orig:FechaSolucion',
    'fechasolucion': 'FechaSolucion',
    'nrocambiofechasolucion': 'cnt:FechaSolucion',
    'foriginalcierre': 'orig:FechaCierre',
    'fechacierre': 'FechaCierre',
    'nrocambiofechacierre': 'cnt:FechaCierre',
}
CAMPOS = ('FechaAnalisis', 'FechaSolucion', 'FechaCierre')

AMBIGUAS = {'FECHA_MAL_FORMADA_NORMALIZADA', 'VARIAS_FECHAS_EN_UNA_ENTRADA', 'SIN_FECHA_NUEVA', 'MISMA_FECHA',
            'FECHA_COMENTARIO_IGUAL_VALOR_ACTUAL', 'REVERSION_A_FECHA_PREVIA', 'SIN_FECHA_EN_COMENTARIO'}

COLUMNAS_CSV = ['IdRevision', 'Codigo', 'Campo', 'CeldaExcel', 'OrdenEnCadena', 'FechaAnterior', 'FechaNueva',
                'FuenteFechaNueva', 'FechaComentario_dT', 'AutorComentario', 'TipoEntrada', 'IdComentarioExcel',
                'TextoComentario', 'TokenFechaOriginal', 'Reconstruido', 'Decision', 'Motivo', 'Ambiguo', 'Flags',
                'Observaciones', 'ValorActualCelda', 'FechaOriginalExcel', 'NroCambioExcel', 'CambiosContadosCelda',
                'TransicionesCelda', 'ChequeoContador']


class ErrorLibro(Exception):
    pass


def norm(texto):
    s = unicodedata.normalize('NFKD', str(texto or ''))
    s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]', '', s.lower())


def ruta_zip(base, destino):
    """Resuelve un Target relativo de un .rels contra la carpeta de su parte."""
    if destino.startswith('/'):
        return destino.lstrip('/')
    partes = base.split('/')[:-1]
    for p in destino.split('/'):
        if p == '..':
            partes.pop()
        elif p not in ('', '.'):
            partes.append(p)
    return '/'.join(partes)


def comentarios_de_hoja(z, hoja):
    """La hoja y su threadedComments, por las relaciones del libro (sin suponer numeros de archivo)."""
    wb = ET.fromstring(z.read('xl/workbook.xml'))
    rels = {e.get('Id'): e.get('Target') for e in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
    for s in wb.find('m:sheets', NS):
        if s.get('name') != hoja:
            continue
        parte = ruta_zip('xl/workbook.xml', rels[s.get('{%s}id' % NS['r'])])
        carpeta, archivo = parte.rsplit('/', 1)
        prels = carpeta + '/_rels/' + archivo + '.rels'
        if prels not in z.namelist():
            raise ErrorLibro('la hoja %s no tiene relaciones: no hay comentarios' % hoja)
        tcs = [ruta_zip(parte, e.get('Target')) for e in ET.fromstring(z.read(prels))
               if e.get('Type', '').endswith('/threadedComment')]
        if len(tcs) != 1:
            raise ErrorLibro('se esperaba 1 archivo de threaded comments para %s, hay %d' % (hoja, len(tcs)))
        return tcs[0]
    raise ErrorLibro('no existe la hoja %s' % hoja)


def personas(z):
    p = 'xl/persons/person.xml'
    if p not in z.namelist():
        return {}
    return {e.get('id'): e.get('displayName') for e in ET.fromstring(z.read(p))}


def columnas(ws):
    encabezado = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    col = {}
    for i, v in enumerate(encabezado, 1):
        papel = ENCABEZADOS.get(norm(v))
        if papel:
            if papel in col:
                raise ErrorLibro('encabezado repetido: %r' % v)
            col[papel] = i
    faltan = [p for p in ENCABEZADOS.values() if p not in col]
    if faltan:
        raise ErrorLibro('faltan columnas en %s: %s' % (HOJA, ', '.join(faltan)))
    return col


def dv(x):
    return x.date() if isinstance(x, datetime.datetime) else x


TOK = re.compile(r'(\d{1,2})\s*/+\s*(\d{1,2})\s*/+\s*(\d{4})')


def tokens(texto):
    out = []
    for m in TOK.finditer(texto):
        raw = m.group(0)
        d, mo, y = map(int, m.groups())
        mal = bool(re.search(r'//|\s', raw))
        try:
            out.append((datetime.date(y, mo, d), raw, mal, ''))
        except ValueError:
            out.append((None, raw, True, 'fecha invalida'))
    return out


def extraer(xlsx):
    z = zipfile.ZipFile(xlsx)
    tc_parte = comentarios_de_hoja(z, HOJA)
    pers = personas(z)
    ws = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)[HOJA]
    col = columnas(ws)
    letra_campo = {get_column_letter(col[c]): c for c in CAMPOS}

    hilos = collections.defaultdict(list)
    for t in ET.fromstring(z.read(tc_parte)).findall('tc:threadedComment', NS):
        ref = t.get('ref')
        if re.sub(r'\d', '', ref) in letra_campo:
            hilos[ref].append(t)
    if not hilos:
        raise ErrorLibro('no hay comentarios en las columnas de fecha')

    filas_necesarias = {coordinate_to_tuple(r)[0] for r in hilos}
    V = {}
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max(filas_necesarias), max_col=max(col.values()),
                                         values_only=True), 1):
        if i in filas_necesarias:
            V[i] = row

    def val(row, papel):
        return row[col[papel] - 1]

    rows = []
    for ref, ts in hilos.items():
        campo = letra_campo[re.sub(r'\d', '', ref)]
        row = V[coordinate_to_tuple(ref)[0]]
        codigo = val(row, 'codigo')
        if not codigo:
            raise ErrorLibro('comentario en %s sin codigo en la fila' % ref)
        codigo = str(codigo).strip()
        cur = dv(val(row, campo))
        orig = dv(val(row, 'orig:' + campo)) if ('orig:' + campo) in col else None
        cnt = val(row, 'cnt:' + campo)
        ts = sorted(ts, key=lambda t: t.get('dT'))
        entradas = []
        for t in ts:
            texto = ''.join(t.find('tc:text', NS).itertext())
            toks = tokens(texto)
            base = dict(dT=t.get('dT'), autor=pers.get(t.get('personId'), t.get('personId')), id=t.get('id'),
                        tipo='respuesta' if t.get('parentId') else 'hilo',
                        texto=texto.strip().replace('\r', '').replace('\n', ' / '))
            if not toks:
                entradas.append(dict(base, fecha=None, raw='', malformed=False, err='sin fecha en texto'))
            for (f, raw, mal, err) in toks:
                entradas.append(dict(base, fecha=f, raw=raw, malformed=mal, err=err, multi=len(toks) > 1))
        cadena = [e for e in entradas if e['fecha']]
        for e in entradas:
            if not e['fecha']:
                rows.append(dict(Codigo=codigo, Campo=campo, Celda=ref, e=e, ant=None, nue=None, fuente='',
                                 decision='NO_CUENTA', motivo=e['err'], flags=['SIN_FECHA_EN_COMENTARIO'], obs=[],
                                 cur=cur, orig=orig, cnt=cnt, orden=''))
        vistas = []
        for i, e in enumerate(cadena):
            ant = e['fecha']
            if i + 1 < len(cadena):
                nue, fuente = cadena[i + 1]['fecha'], 'SIGUIENTE_COMENTARIO'
            elif isinstance(cur, datetime.date):
                nue, fuente = cur, 'VALOR_ACTUAL_CELDA'
            else:
                nue, fuente = None, 'SIN_VALOR_ACTUAL'
            flags, obs = [], []
            if e['malformed']:
                flags.append('FECHA_MAL_FORMADA_NORMALIZADA')
            if e.get('multi'):
                flags.append('VARIAS_FECHAS_EN_UNA_ENTRADA')
            if nue is None:
                flags.append('SIN_FECHA_NUEVA')
            elif nue == ant:
                flags.append('MISMA_FECHA')
            if ant == cur:
                flags.append('FECHA_COMENTARIO_IGUAL_VALOR_ACTUAL')
            if nue is not None and nue != ant and nue in vistas:
                flags.append('REVERSION_A_FECHA_PREVIA')
            if i == 0 and orig and orig != ant:
                obs.append('F. Original distinta de la 1a fecha comentada: intermedios sin evidencia')
            if campo != 'FechaCierre' and nue is not None and nue < ant:
                obs.append('fecha adelantada (cuenta igual)')
            if campo == 'FechaCierre' and nue is not None and nue == datetime.date.fromisoformat(e['dT'][:10]):
                obs.append('cierre: FechaNueva = dia del comentario')
            vistas.append(ant)
            if nue is None:
                dec, mot = 'NO_CUENTA', 'sin fecha nueva: celda actual vacia o no es fecha'
            elif nue == ant:
                dec, mot = ('EXCLUIDA_CIERRE' if campo == 'FechaCierre' else 'NO_CUENTA'), 'misma fecha: no hay cambio'
            elif campo == 'FechaCierre' and nue < ant:
                dec, mot = 'EXCLUIDA_CIERRE', 'cierre: FechaNueva < FechaAnterior, no es extension'
            else:
                dec, mot = 'CUENTA', ('cierre: extension' if campo == 'FechaCierre' else 'cambio de fecha soportado')
            rows.append(dict(Codigo=codigo, Campo=campo, Celda=ref, e=e, ant=ant, nue=nue, fuente=fuente,
                             decision=dec, motivo=mot, flags=flags, obs=obs, cur=cur, orig=orig, cnt=cnt, orden=i + 1))
    return rows, len(hilos)


def chequeo(percell, celda, cnt):
    if cnt in (None, ''):
        return 'SIN_CONTADOR'
    try:
        c = int(cnt)
    except (TypeError, ValueError):
        return 'CONTADOR_NO_NUMERICO'
    n = percell[celda][0]
    return 'CUADRA' if c == n else ('CONTADOR_MAYOR' if c > n else 'CONTADOR_MENOR')


def a_csv(rows):
    percell = collections.defaultdict(lambda: [0, 0])
    for x in rows:
        if x['ant'] is not None:
            percell[x['Celda']][1] += 1
        if x['decision'] == 'CUENTA':
            percell[x['Celda']][0] += 1
    rows.sort(key=lambda x: (x['Codigo'], x['Campo'], x['e']['dT'], x['e']['raw']))
    out = []
    for i, x in enumerate(rows, 1):
        e = x['e']
        amb = bool(AMBIGUAS & set(x['flags']))
        out.append(dict(zip(COLUMNAS_CSV, [
            i, x['Codigo'], x['Campo'], x['Celda'], x['orden'], x['ant'] or '', x['nue'] or '', x['fuente'],
            e['dT'], e['autor'], e['tipo'], e['id'], e['texto'], e['raw'], 1, x['decision'], x['motivo'],
            'SI' if amb else 'NO', ';'.join(x['flags']), ';'.join(x['obs']),
            x['cur'] if x['cur'] is not None else '', x['orig'] or '', '' if x['cnt'] is None else x['cnt'],
            percell[x['Celda']][0], percell[x['Celda']][1], chequeo(percell, x['Celda'], x['cnt'])])))
    return out


def validar_cadenas(filas):
    """Lo que el INSERT exige: transicion real, cadenas sin saltos, ultima = valor actual."""
    problemas = []
    porcampo = collections.defaultdict(list)
    for r in filas:
        if r['Decision'] != 'CUENTA':
            continue
        if not r['FechaAnterior'] or not r['FechaNueva'] or str(r['FechaAnterior']) == str(r['FechaNueva']):
            problemas.append('transicion invalida en IdRevision %s' % r['IdRevision'])
        porcampo[(r['Codigo'], r['Campo'])].append(r)
    for k, v in porcampo.items():
        v.sort(key=lambda r: int(r['OrdenEnCadena']))
        for a, b in zip(v, v[1:]):
            if str(a['FechaNueva']) != str(b['FechaAnterior']):
                problemas.append('salto en %s %s entre ordenes %s y %s' % (k + (a['OrdenEnCadena'], b['OrdenEnCadena'])))
        if str(v[-1]['FechaNueva']) != str(v[-1]['ValorActualCelda']):
            problemas.append('%s %s: la ultima fecha nueva no es el valor actual' % k)
    return problemas, len(porcampo)


def diferencias(nuevas, ruta_anterior):
    with open(ruta_anterior, encoding='utf-8-sig') as f:
        viejas = list(csv.DictReader(f))
    llave = lambda r: (r['IdComentarioExcel'], r['TokenFechaOriginal'], r['CeldaExcel'])
    comp = ('Codigo', 'Campo', 'FechaAnterior', 'FechaNueva', 'Decision', 'Flags')
    V = {llave(r): r for r in viejas}
    N = {llave(r): r for r in nuevas}
    out = []
    for k, r in N.items():
        if k not in V:
            out.append(dict(Cambio='NUEVA', **{c: r[c] for c in COLUMNAS_CSV}))
        else:
            dif = [c for c in comp if str(V[k][c]) != str(r[c])]
            if dif:
                out.append(dict(Cambio='CAMBIO:' + ','.join(dif), **{c: r[c] for c in COLUMNAS_CSV}))
    for k, r in V.items():
        if k not in N:
            out.append(dict(Cambio='YA_NO_ESTA', **{c: r[c] for c in COLUMNAS_CSV}))
    return out


def q(v):
    v = '' if v is None else str(v)
    return 'NULL' if v == '' else "N'" + v.replace("'", "''") + "'"


SQL_CABECERA = r"""/* =====================================================================
   Importacion del historial de fechas desde los comentarios del Excel
   Generado por tools/historial_fechas/extraer_historial_fechas.py
   Fuente: {fuente}
   Generado: {generado}   Filas CUENTA: {n_cuenta}

   POR OMISION NO CAMBIA NADA. Las secciones 0-3 solo leen. La 4 inserta
   dentro de una transaccion y termina en ROLLBACK salvo que se pongan las
   DOS variables de abajo:
     @ModoFechaRegistro  'COMENTARIO'  -> FechaRegistro = dT del comentario
                                         (zona horaria del dT SIN verificar)
                         'IMPORTACION' -> FechaRegistro = ahora (default)
                         NULL          -> la seccion 4 se niega a insertar
     @Confirmar          0 -> ROLLBACK (ensayo)   1 -> COMMIT

   Cada fila entra como Operacion 'U', Origen 'NO_DECLARADO', Usuario y
   SolicitudId NULL; LoginBD = quien corre el script (default de la tabla).
   NO crea linea base ('B') ni triggers. Se niega a insertar si ya existe
   el trigger de captura en dbo.Problem: despues de eso el historial ya no
   es solo de esta importacion.

   Se puede volver a correr con un Excel mas nuevo: por cada Codigo+Campo
   compara lo que ya hay con la cadena del Excel, en orden. Si lo que hay es
   el principio de la cadena, inserta solo lo que falta; si no coincide, no
   toca ese campo y lo lista en 3c.

   Un ensayo (ROLLBACK) consume valores de IDENTITY: los IdEvento del COMMIT
   real empiezan despues. No afecta el orden de las cadenas.
   ===================================================================== */
SET NOCOUNT ON;
SET XACT_ABORT ON;

DECLARE @ModoFechaRegistro varchar(11) = NULL;   -- 'COMENTARIO' | 'IMPORTACION'
DECLARE @Confirmar bit = 0;                      -- 1 = COMMIT

/* ---------- 0. Destino (solo lectura) ---------- */
SELECT Seccion = '0 destino',
       TablaExiste    = CASE WHEN OBJECT_ID(N'dbo.ProblemFechaEvento', N'U') IS NULL THEN 0 ELSE 1 END,
       TriggerCaptura = (SELECT COUNT(*) FROM sys.triggers WHERE parent_id = OBJECT_ID(N'dbo.Problem')),
       FilasHoy       = (SELECT COUNT_BIG(*) FROM dbo.ProblemFechaEvento);

/* ---------- 1. Staging (temporal) ---------- */
IF OBJECT_ID('tempdb..#R') IS NOT NULL DROP TABLE #R;
CREATE TABLE #R (
    IdRevision         int           NOT NULL PRIMARY KEY,
    Codigo             nvarchar(100) COLLATE SQL_Latin1_General_CP1_CI_AS NOT NULL,
    Campo              varchar(13)   NOT NULL,
    CeldaExcel         varchar(10)   NOT NULL,
    OrdenEnCadena      int           NOT NULL,
    FechaAnterior      date          NOT NULL,
    FechaNueva         date          NOT NULL,
    FechaComentario    datetime2(3)  NOT NULL,   -- dT del comentario (zona sin verificar)
    IdComentarioExcel  varchar(40)   NOT NULL,
    Ambiguo            varchar(2)    NOT NULL,
    ValorActualExcel   date          NULL
);

INSERT INTO #R VALUES
"""

SQL_CUERPO = r""";

-- Orden dentro de la cadena de CUENTA (1, 2, ...) por Codigo+Campo.
IF OBJECT_ID('tempdb..#C') IS NOT NULL DROP TABLE #C;
SELECT r.*, N = ROW_NUMBER() OVER (PARTITION BY r.Codigo, r.Campo ORDER BY r.OrdenEnCadena)
INTO #C FROM #R AS r;

-- Lo que ya hay en la tabla de esa misma clase de fila, en orden de IdEvento.
IF OBJECT_ID('tempdb..#E') IS NOT NULL DROP TABLE #E;
SELECT e.Codigo, e.Campo, e.FechaAnterior, e.FechaNueva,
       N = ROW_NUMBER() OVER (PARTITION BY e.Codigo, e.Campo ORDER BY e.IdEvento)
INTO #E
FROM dbo.ProblemFechaEvento AS e
WHERE e.Operacion = 'U' AND e.Origen = 'NO_DECLARADO';

/* ---------- 2. Vista previa (solo lectura) ---------- */
-- 2a. Totales
SELECT Seccion = '2a totales', Campo, Ambiguo, Filas = COUNT(*) FROM #R GROUP BY Campo, Ambiguo ORDER BY Campo, Ambiguo;

-- 2b. Codigos que no existen en dbo.Problem (la FK los rechazaria): no se insertan.
SELECT Seccion = '2b sin Problem', r.Codigo, Filas = COUNT(*)
FROM #R AS r WHERE NOT EXISTS (SELECT 1 FROM dbo.Problem AS p WHERE p.Codigo = r.Codigo)
GROUP BY r.Codigo;

-- 2c. Valor actual: Excel vs base. Si difieren, la cadena no termina donde esta la base hoy.
SELECT Seccion = '2c valor distinto', x.Codigo, x.Campo, x.ValorActualExcel, x.ValorBD
FROM (SELECT DISTINCT r.Codigo, r.Campo, r.ValorActualExcel,
             ValorBD = CONVERT(date, CASE r.Campo WHEN 'FechaAnalisis' THEN p.FechaAnalisis
                                                  WHEN 'FechaSolucion' THEN p.FechaSolucion
                                                  ELSE p.FechaCierre END)
      FROM #R AS r JOIN dbo.Problem AS p ON p.Codigo = r.Codigo) AS x
WHERE EXISTS (SELECT x.ValorActualExcel EXCEPT SELECT x.ValorBD);

/* ---------- 3. Contra lo que ya hay ---------- */
-- Por Codigo+Campo: cuantas hay (m) y si son el principio exacto de la cadena del Excel.
IF OBJECT_ID('tempdb..#K') IS NOT NULL DROP TABLE #K;
SELECT k.Codigo, k.Campo,
       Existentes = (SELECT COUNT(*) FROM #E AS e WHERE e.Codigo = k.Codigo AND e.Campo = k.Campo),
       EnExcel    = (SELECT COUNT(*) FROM #C AS c WHERE c.Codigo = k.Codigo AND c.Campo = k.Campo),
       Conflicto  = CASE WHEN EXISTS (
                        SELECT 1 FROM #E AS e
                        LEFT JOIN #C AS c ON c.Codigo = e.Codigo AND c.Campo = e.Campo AND c.N = e.N
                        WHERE e.Codigo = k.Codigo AND e.Campo = k.Campo
                          AND (c.N IS NULL OR c.FechaAnterior <> e.FechaAnterior OR c.FechaNueva <> e.FechaNueva))
                    THEN 1 ELSE 0 END
INTO #K
FROM (SELECT DISTINCT Codigo, Campo FROM #C UNION SELECT DISTINCT Codigo, Campo FROM #E) AS k;

-- 3a. Resumen
SELECT Seccion = '3a resumen',
       CamposExcel     = SUM(CASE WHEN EnExcel > 0 THEN 1 ELSE 0 END),
       YaCompletos     = SUM(CASE WHEN Conflicto = 0 AND Existentes = EnExcel AND EnExcel > 0 THEN 1 ELSE 0 END),
       ConFilasNuevas  = SUM(CASE WHEN Conflicto = 0 AND EnExcel > Existentes THEN 1 ELSE 0 END),
       EnConflicto     = SUM(Conflicto)
FROM #K;

-- 3b. Lo que se insertaria
IF OBJECT_ID('tempdb..#I') IS NOT NULL DROP TABLE #I;
SELECT c.*
INTO #I
FROM #C AS c
JOIN #K AS k ON k.Codigo = c.Codigo AND k.Campo = c.Campo
WHERE k.Conflicto = 0
  AND c.N > k.Existentes
  AND EXISTS (SELECT 1 FROM dbo.Problem AS p WHERE p.Codigo = c.Codigo);
SELECT Seccion = '3b a insertar', Filas = COUNT(*), Ambiguas = SUM(CASE WHEN Ambiguo = 'SI' THEN 1 ELSE 0 END) FROM #I;
SELECT Seccion = '3b detalle', Codigo, Campo, N, FechaAnterior, FechaNueva, FechaComentario, Ambiguo, CeldaExcel
FROM #I ORDER BY Codigo, Campo, N;

-- 3c. Campos en conflicto: lo que hay no es el principio de la cadena. No se tocan.
SELECT Seccion = '3c conflicto', k.Codigo, k.Campo, k.Existentes, k.EnExcel FROM #K AS k WHERE k.Conflicto = 1;

/* ---------- 4. INSERT (transaccion; ROLLBACK salvo @Confirmar = 1) ---------- */
IF @ModoFechaRegistro IS NULL OR @ModoFechaRegistro NOT IN ('COMENTARIO', 'IMPORTACION')
BEGIN
    SELECT Seccion = '4 omitida', Motivo = 'Falta @ModoFechaRegistro (COMENTARIO o IMPORTACION). Nada se inserto.';
    RETURN;
END;
IF EXISTS (SELECT 1 FROM sys.triggers WHERE parent_id = OBJECT_ID(N'dbo.Problem'))
BEGIN
    SELECT Seccion = '4 omitida', Motivo = 'Ya hay trigger en dbo.Problem: el backfill ya no aplica. Nada se inserto.';
    RETURN;
END;

BEGIN TRAN;

DECLARE @antes bigint = (SELECT COUNT_BIG(*) FROM dbo.ProblemFechaEvento WITH (TABLOCKX, HOLDLOCK));

-- ORDER BY en INSERT...SELECT fija el orden de IdEvento: la cadena queda en orden.
IF @ModoFechaRegistro = 'COMENTARIO'
    INSERT INTO dbo.ProblemFechaEvento (Codigo, Campo, FechaAnterior, FechaNueva, Operacion, Origen, FechaRegistro)
    SELECT Codigo, Campo, FechaAnterior, FechaNueva, 'U', 'NO_DECLARADO', FechaComentario
    FROM #I ORDER BY Codigo, Campo, N;
ELSE
    INSERT INTO dbo.ProblemFechaEvento (Codigo, Campo, FechaAnterior, FechaNueva, Operacion, Origen)
    SELECT Codigo, Campo, FechaAnterior, FechaNueva, 'U', 'NO_DECLARADO'
    FROM #I ORDER BY Codigo, Campo, N;

DECLARE @insertadas int = @@ROWCOUNT;
DECLARE @esperadas int = (SELECT COUNT(*) FROM #I);
DECLARE @despues bigint = (SELECT COUNT_BIG(*) FROM dbo.ProblemFechaEvento);

SELECT Seccion = '4 resultado', Modo = @ModoFechaRegistro, Antes = @antes, Insertadas = @insertadas,
       Esperadas = @esperadas, Despues = @despues,
       Accion = CASE WHEN @Confirmar = 1 AND @insertadas = @esperadas THEN 'COMMIT' ELSE 'ROLLBACK' END;

IF @Confirmar = 1 AND @insertadas = @esperadas
    COMMIT;
ELSE
    ROLLBACK;
"""


def a_sql(filas, fuente):
    cuenta = [r for r in filas if r['Decision'] == 'CUENTA']
    vals = []
    for r in cuenta:
        dt = r['FechaComentario_dT']
        vals.append('(%s)' % ', '.join([
            str(r['IdRevision']), q(r['Codigo']), q(r['Campo']), q(r['CeldaExcel']), str(int(r['OrdenEnCadena'])),
            q(r['FechaAnterior']), q(r['FechaNueva']), q(dt), q(r['IdComentarioExcel']), q(r['Ambiguo']),
            q(r['ValorActualCelda'])]))
    cab = SQL_CABECERA.format(fuente=os.path.basename(fuente).replace('*/', ''),
                              generado=datetime.datetime.now().strftime('%Y-%m-%d %H:%M'), n_cuenta=len(cuenta))
    return cab + ',\n'.join(vals) + SQL_CUERPO, len(cuenta)


def escribir_csv(ruta, filas, columnas):
    with open(ruta, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=columnas)
        w.writeheader()
        w.writerows(filas)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('libro')
    ap.add_argument('salida')
    ap.add_argument('--anterior', help='Revision.csv anterior, para listar diferencias')
    a = ap.parse_args()
    os.makedirs(a.salida, exist_ok=True)

    try:
        rows, n_celdas = extraer(a.libro)
    except ErrorLibro as e:
        print('ERROR:', e)
        return 2
    filas = a_csv(rows)

    problemas, n_cadenas = validar_cadenas(filas)
    if problemas:
        print('ERROR: la revision no cumple lo que exige la tabla; no se genera SQL:')
        for p in problemas:
            print('  -', p)
        escribir_csv(os.path.join(a.salida, 'ProblemFechaEvento_Revision.csv'), filas, COLUMNAS_CSV)
        return 3

    escribir_csv(os.path.join(a.salida, 'ProblemFechaEvento_Revision.csv'), filas, COLUMNAS_CSV)
    sql, n = a_sql(filas, a.libro)
    with open(os.path.join(a.salida, 'ProblemFechaEvento_Importar.sql'), 'w', encoding='utf-8-sig') as f:
        f.write(sql)

    S = collections.Counter((r['Campo'], r['Decision']) for r in filas)
    amb = sum(1 for r in filas if r['Decision'] == 'CUENTA' and r['Ambiguo'] == 'SI')
    print('celdas con comentarios:', n_celdas, ' filas:', len(filas),
          ' codigos:', len({r['Codigo'] for r in filas}))
    for k in sorted(S):
        print('  %-14s %-16s %d' % (k[0], k[1], S[k]))
    print('CUENTA:', n, ' en', n_cadenas, 'cadenas;  CUENTA ambiguas:', amb)

    if a.anterior:
        dif = diferencias(filas, a.anterior)
        escribir_csv(os.path.join(a.salida, 'ProblemFechaEvento_Diferencias.csv'), dif, ['Cambio'] + COLUMNAS_CSV)
        print('diferencias contra la anterior:', collections.Counter(d['Cambio'].split(':')[0] for d in dif))
    return 0


if __name__ == '__main__':
    sys.exit(main())
